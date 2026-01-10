from docx import Document

# 预加载中文字符范围
chinese_chars = set(chr(i) for i in range(0x4e00, 0x9fff + 1))

def is_chinese(text):
    """
    判断字符串是否完全由中文字符组成。
    :param text: 输入的字符串
    :return: 如果完全由中文字符组成返回 True，否则返回 False
    """
    return all(char in chinese_chars for char in text)

def separate_non_chinese(input_file, output_docx_file, index_file):
    """
    提取第二列不完全由中文字符组成的行，并保存为 .docx 文件。
    :param input_file: 输入文件路径
    :param output_docx_file: 输出的 .docx 文件路径
    :param index_file: 行号索引文件路径
    """
    doc = Document()  # 创建一个新的 Word 文档
    with open(input_file, 'r', encoding='utf-8') as infile, open(index_file, 'w', encoding='utf-8') as idxfile:
        line_number = 0
        for line in infile:
            line_number += 1
            parts = line.strip().split('\t')  # 按制表符分割
            if len(parts) >= 2 and not parts[0]==parts[1] and not is_chinese(parts[1]):  # 判断第二列
                doc.add_paragraph(parts[0])  # 将第一列添加到 Word 文档
                # doc.add_paragraph(parts[1])  # 将第二列添加到 Word 文档
                idxfile.write(f"{line_number}\n")  # 记录行号

    # 保存 Word 文档
    doc.save(output_docx_file)

# 示例调用
language_pairs = [
    "en-af", "en-ar", "en-bg", "en-bn", "en-bs", "en-ca", "en-cs", "en-da", "en-de",
    "en-el", "en-en", "en-es", "en-et", "en-fa", "en-fi", "en-fr", "en-he", "en-hi",
    "en-hr", "en-hu", "en-id", "en-it", "en-ja", "en-ko", "en-lt", "en-lv", "en-mk",
    "en-ms", "en-nl", "en-no", "en-pl", "en-pt", "en-ro", "en-ru", "en-sk", "en-sl",
    "en-sq", "en-sv", "en-ta", "en-th", "en-tl", "en-tr", "en-uk", "en-vi", "en-zh"
]
# pair = input()
for pair in language_pairs:
    input_file = rf'F:\PythonProject\tranlate\output\translated_{pair}.txt'  # 输入文件名
    output_docx_file = rf"E:\onedrive\桌面\lan\{pair}.docx"  # 输出文件目录
    index_file = rf"E:\onedrive\桌面\lan\{pair}_index.txt"  # 行号索引文件路径

    separate_non_chinese(input_file, output_docx_file, index_file)
    print(f"不完全由中文字符组成的行已分离并保存到: {output_docx_file}")
    print(f"行号索引已保存到: {index_file}")